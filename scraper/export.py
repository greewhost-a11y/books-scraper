"""Write scraped rows to CSV, JSON or a formatted Excel workbook."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATS = ("csv", "json", "xlsx")


def export(rows: list[dict], path: Path) -> Path:
    fmt = path.suffix.lstrip(".").lower()
    if fmt not in FORMATS:
        raise ValueError(f"Unsupported format {fmt!r}; use one of {FORMATS}")
    path.parent.mkdir(parents=True, exist_ok=True)
    {"csv": _csv, "json": _json, "xlsx": _xlsx}[fmt](rows, path)
    return path


def _csv(rows: list[dict], path: Path) -> None:
    # utf-8-sig so Excel opens non-ASCII text correctly.
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        if not rows:
            return
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _json(rows: list[dict], path: Path) -> None:
    path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


def _xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Books"
    if rows:
        headers = list(rows[0])
        ws.append(headers)
        for row in rows:
            ws.append([row[h] for h in headers])

        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill

        for idx, header in enumerate(headers, start=1):
            longest = max(len(str(r[header])) for r in rows)
            ws.column_dimensions[get_column_letter(idx)].width = min(max(len(header), longest) + 2, 60)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
